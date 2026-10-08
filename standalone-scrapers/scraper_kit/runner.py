"""公共运行逻辑：命令行(run_demo.py)和网页(app.py)共用。"""
from __future__ import annotations

import time
from typing import Any

from openpyxl import Workbook

from .registry import discover_scrapers

# 每个来源: (中文表名, [(原字段, 中文列名), ...])。不在列表里的字段(含重复字段 avg_price / deals_count)不输出。
# 链家的 avg_price == unit_price_yuan_sqm、deals_count == follow_count，故去重。
COLUMNS = {
    "lianjia_deals": ("链家二手房挂牌", [
        ("city", "城市"), ("district", "区域"), ("xiaoqu", "小区"), ("period", "期间"),
        ("unit_price_yuan_sqm", "挂牌单价(元/平方米)"), ("total_price_wan", "挂牌总价(万元)"),
        ("follow_count", "关注人数(人)"), ("title", "房源标题"),
        ("source", "数据来源"), ("is_fallback", "是否兜底数据"),
    ]),
    "nbs_house_price": ("统计局70城房价指数", [
        ("city", "城市"), ("period", "期间"),
        ("new_home_index_yoy", "新建商品住宅同比(%)"), ("new_home_index_mom", "新建商品住宅环比(%)"),
        ("second_hand_index_yoy", "二手住宅同比(%)"), ("second_hand_index_mom", "二手住宅环比(%)"),
        ("source", "数据来源"), ("is_fallback", "是否兜底数据"),
    ]),
    "policy_crawler": ("房地产政策", [
        ("policy_id", "政策编号"), ("title", "标题"), ("publish_date", "发布日期"),
        ("city", "城市"), ("level", "级别"), ("content", "内容摘要"), ("source_url", "原文链接"),
        ("source", "数据来源"), ("is_fallback", "是否兜底数据"),
    ]),
}


async def run_scraper(s):
    """复刻 BaseScraper.run，但保留数据行，并明确标出是否为兜底数据。"""
    t0 = time.perf_counter()
    try:
        rows = s.validate(s.parse(await s.fetch()))
        status, err, fb = "实时数据", "", False
    except Exception as exc:  # noqa: BLE001
        rows = s.validate(s.fallback())
        status, err, fb = "兜底(非真实)数据", f"{type(exc).__name__}: {exc}", True
    return {"s": s, "rows": rows, "status": status, "error": err, "fallback": fb,
            "ms": int((time.perf_counter() - t0) * 1000)}




def build_workbook(results: list[dict[str, Any]]) -> Workbook:
    """把运行结果写成 Excel：一个"运行摘要"页 + 每个来源一页，表头为中文并带单位。"""
    wb = Workbook()
    wb.remove(wb.active)
    summary = wb.create_sheet("运行摘要")
    summary.append(["来源", "状态", "行数", "耗时(毫秒)", "失败原因"])
    for r in results:
        s = r["s"]
        summary.append([s.name, r["status"], len(r["rows"]), r["ms"], r["error"]])
        sheet, cols = COLUMNS.get(s.source_id, (s.source_id[:31], None))
        ws = wb.create_sheet(sheet)
        if r["rows"]:
            if cols is None:
                cols = [(k, k) for k in r["rows"][0].keys()]
            ws.append([zh for _, zh in cols])
            for row in r["rows"]:
                ws.append([row.get(k) for k, _ in cols])
            for i, (_, zh) in enumerate(cols, 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(zh) * 2 + 2)
    return wb


def to_records(result: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    """返回 (中文列名, 行数据)，供网页表格展示。"""
    _, cols = COLUMNS.get(result["s"].source_id, ("", None))
    rows = result["rows"]
    if cols is None and rows:
        cols = [(k, k) for k in rows[0].keys()]
    cols = cols or []
    return [zh for _, zh in cols], [[row.get(k) for k, _ in cols] for row in rows]
